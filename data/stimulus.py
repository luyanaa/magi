"""Stimulus timing metadata and the salt-stimulus waveform.

The salt drive used with the ``cleandata_smoothened2`` recordings is defined by
``generatesalt.m`` in the gKDR-GMM release (Toyoshima et al., PLOS Comput Biol
2024, doi:10.1371/journal.pcbi.1011848)::

    x = (1:n) - startframe
    y = (abs(sin(x/period*pi))).^0.25 .* sign(sin(x/period*pi))
    y(1:floor(startframe)) = 0

``startframe`` and ``period`` are in FRAMES, so reproducing the drive requires
the per-sample frame rate (``stimulation_timing.xlsx``). ``period`` is the
interval between NaCl concentration changes (the paper's 30 s half-period), not
the full cycle.

The per-sample table also carries the animal id, the anaesthesia flag and the
recording duration, which the ingest writes into the manifest so that subject
grouping and condition metadata survive into the ladder.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union

import numpy as np

__all__ = [
    "SampleTiming",
    "read_stimulation_timing",
    "salt_waveform",
    "trial_windows",
]


@dataclass(frozen=True)
class SampleTiming:
    """One row of the stimulation-timing table (frames are per-sample)."""

    sample: int
    animal: int
    name: str
    anesthesia: int
    fps: float
    first_stimulus_frame: float
    switch_period_frames: float
    cycles: float
    total_duration_s: float

    @property
    def dt_s(self) -> float:
        return 1.0 / float(self.fps)

    @property
    def first_stimulus_s(self) -> float:
        return float(self.first_stimulus_frame) / float(self.fps)

    @property
    def switch_period_s(self) -> float:
        return float(self.switch_period_frames) / float(self.fps)


def _as_float(value) -> float:
    return float(value) if value is not None else float("nan")


def read_stimulation_timing(path: Union[str, Path]) -> Dict[int, SampleTiming]:
    """Parse a stimulation-timing table (xlsx or csv), keyed by sample index.

    Expected columns (the gKDR-GMM ``stimulation_timing.xlsx`` layout):
    ``#sample, #animal, name, anesthesia, frames/sec, frame number of first
    stimuli (...), every N frames NaCl concentration change, number of
    stimulation cycles, total duration (sec)``. Column matching is prefix
    based and case/space tolerant so minor header edits do not break parsing.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"stimulation timing table not found: {path}")
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - env dependent
            raise ImportError(
                "reading a stimulation-timing .xlsx requires openpyxl: "
                "pip install openpyxl") from exc
        wb = openpyxl.load_workbook(path, data_only=True)
        rows = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    elif path.suffix.lower() in (".csv", ".tsv", ".txt"):
        import csv as _csv
        delim = "\t" if path.suffix.lower() == ".tsv" else ","
        with open(path, newline="") as fh:
            rows = [list(r) for r in _csv.reader(fh, delimiter=delim)]
    else:
        raise ValueError(f"unsupported timing table format: {path.suffix}")

    header_idx = None
    for i, row in enumerate(rows):
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        if any(c.startswith("#sample") or c == "sample" for c in cells):
            header_idx = i
            break
    if header_idx is None:
        raise ValueError(f"{path}: no '#sample' header row found")

    def find(header: List[str], *prefixes: str) -> int:
        for i, cell in enumerate(header):
            for p in prefixes:
                if cell.startswith(p):
                    return i
        raise ValueError(
            f"{path}: missing column {prefixes[0]!r}; header={header}")

    header = [str(c).strip().lower() if c is not None else "" for c in rows[header_idx]]
    i_sample = find(header, "#sample", "sample")
    i_animal = find(header, "#animal", "animal")
    i_name = find(header, "name")
    i_anes = find(header, "anesth")
    i_fps = find(header, "frames/sec", "frame/sec", "fps")
    i_first = find(header, "frame number of first", "first stimuli", "first stimulus")
    i_period = find(header, "every n frames", "every")
    i_cycles = find(header, "number of stimulation cycles", "cycles")
    i_dur = find(header, "total duration", "duration")

    out: Dict[int, SampleTiming] = {}
    for row in rows[header_idx + 1:]:
        if row is None or len(row) <= i_sample or row[i_sample] in (None, ""):
            continue
        sample = int(float(row[i_sample]))
        out[sample] = SampleTiming(
            sample=sample,
            animal=int(float(row[i_animal])),
            name=str(row[i_name]).strip(),
            anesthesia=int(float(row[i_anes])),
            fps=_as_float(row[i_fps]),
            first_stimulus_frame=_as_float(row[i_first]),
            switch_period_frames=_as_float(row[i_period]),
            cycles=_as_float(row[i_cycles]),
            total_duration_s=_as_float(row[i_dur]),
        )
    if not out:
        raise ValueError(f"{path}: no sample rows parsed")
    bad = {s: t.fps for s, t in out.items() if not np.isfinite(t.fps) or t.fps <= 0}
    if bad:
        raise ValueError(f"{path}: non-positive frames/sec for samples {sorted(bad)}")
    return out


def salt_waveform(n_frames: int,
                  start_frame: float,
                  period_frames: float) -> np.ndarray:
    """Salt drive for ``n_frames`` samples, matching gKDR-GMM ``generatesalt.m``.

    ``start_frame``/``period_frames`` are in frames of the same recording.
    The pre-stimulus span is exactly zero, and the drive has the sub-linear
    (|.|^0.25) shape and alternating sign of the reference implementation.
    """
    if n_frames <= 0:
        raise ValueError("n_frames must be positive")
    if not np.isfinite(start_frame) or start_frame < 0:
        raise ValueError("start_frame must be finite and non-negative")
    if not np.isfinite(period_frames) or period_frames <= 0:
        raise ValueError("period_frames must be finite and positive")
    x = np.arange(1, n_frames + 1, dtype=np.float64) - float(start_frame)
    phase = np.sin(x / float(period_frames) * np.pi)
    y = np.sign(phase) * np.abs(phase) ** 0.25
    y[: int(np.floor(start_frame))] = 0.0
    return y.astype(np.float32)


def trial_windows(n_frames: int,
                  start_frame: float,
                  period_frames: float) -> List[Tuple[int, int]]:
    """Constant-concentration windows ``[lo, hi)`` implied by the switch period.

    One window per half-cycle (25 mM / 50 mM), starting after the pre-stimulus
    span. Windows tile the post-stimulus span exactly (no overlap, no gap) and
    are clipped to ``n_frames``.
    """
    if n_frames <= 0:
        raise ValueError("n_frames must be positive")
    start = float(start_frame)
    period = float(period_frames)
    if not np.isfinite(start) or start < 0:
        raise ValueError("start_frame must be finite and non-negative")
    if not np.isfinite(period) or period <= 0:
        raise ValueError("period_frames must be finite and positive")
    edges = [int(np.floor(start + k * period))
             for k in range(int(np.ceil((n_frames - start) / period)) + 1)]
    windows: List[Tuple[int, int]] = []
    for lo, hi in zip(edges[:-1], edges[1:]):
        lo = max(0, min(int(n_frames), lo))
        hi = max(0, min(int(n_frames), hi))
        if hi > lo:
            windows.append((lo, hi))
    return windows
